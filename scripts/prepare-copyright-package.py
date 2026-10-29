"""Prepare registration drafts from a frozen committed source snapshot.

No applicant identity is guessed. Outputs are private under ignored output/.
Run with the bundled Python runtime, after capturing actual UI screenshots.
"""
from pathlib import Path
import hashlib
import json
import re
import subprocess
import zipfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/copyright-v235'
SNAP = OUT / 'snapshot'
TITLE = '灰线像素前线游戏软件'
VERSION = 'V1.0'
HEADER = TITLE + ' ' + VERSION
FONT = '/System/Library/Fonts/Supplemental/Arial Unicode.ttf'
pdfmetrics.registerFont(TTFont('CJK', FONT))
LAW = 'https://www.ncac.gov.cn/xxfb/flfg/bmgz/202410/t20241015_869486.html'
NOTICE = 'https://www.ccct.net.cn/html/bqzx/2023/0531/4368.html'
PORTAL = 'https://register.ccopyright.com.cn/registration.html#/registerSoft'


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def dependencies(entries):
    seen = set()
    def visit(p):
        p = p.resolve()
        if p in seen or p.suffix not in ('.ts', '.tsx'):
            return
        assert p.is_relative_to(SNAP)
        seen.add(p)
        for spec in re.findall(r'''(?:from\s*|import\s*\(\s*|import\s*)['"]([^'"]+)['"]''', p.read_text()):
            target = SNAP / spec[2:] if spec.startswith('@/') else p.parent / spec if spec.startswith('.') else None
            if target is None:
                continue
            choices = [target] if target.suffix else [target.with_suffix('.ts'), target.with_suffix('.tsx'), target / 'index.ts', target / 'index.tsx']
            for f in choices:
                if f.exists() and f.suffix in ('.ts', '.tsx'):
                    visit(f)
                    break
    for e in entries:
        visit(e)
    # These two files contain generated image resource metadata, not authored
    # program logic. They remain in the frozen snapshot and exclusion evidence.
    excluded = {p for p in seen if p.name in ('battle-art-baked-data.ts', 'card-image-data.ts')}
    included = seen - excluded
    return entries + sorted(included - set(entries), key=lambda p: p.relative_to(SNAP).as_posix()), sorted(excluded)


def wrap(text, size, width):
    lines, buf, length = [], '', 0
    for c in text.expandtabs(2):
        w = pdfmetrics.stringWidth(c, 'CJK', size)
        if length + w > width and buf:
            lines.append(buf)
            buf, length = '', 0
        buf += c
        length += w
    return lines + [buf]


def source_pdf():
    files, excluded = dependencies([SNAP / 'minigame/src/main.ts', SNAP / 'static-app/main.tsx'])
    records = [(p.relative_to(SNAP).as_posix(), n, s) for p in files for n, s in enumerate(p.read_text().splitlines(), 1) if s.strip()]
    assert len(records) > 3000
    selected = records[:1500] + records[-1500:]
    out = OUT / '01-source-code.pdf'
    pdf = canvas.Canvas(str(out), pagesize=A4, pageCompression=1)
    pdf.setTitle(HEADER + ' 源程序鉴别材料')
    W, H = A4
    evidence = []
    text_pages = []
    for i in range(60):
        rows = selected[i*50:(i+1)*50]
        size = 9.0
        while True:
            wrapped = [wrap(r[2], size, W-101) for r in rows]
            count = sum(map(len, wrapped))
            leading = min(13.5, (H-106)/count)
            if leading >= size*1.15:
                break
            size -= .25
            assert size >= 7, f'Page {i+1} would be unreadable'
        pdf.setFillColorRGB(0, 0, 0)
        pdf.setFont('CJK', 10)
        pdf.drawString(36, H-27, HEADER)
        pdf.drawRightString(W-36, H-27, f'第 {i+1} 页')
        pdf.setFont('CJK', 7)
        pdf.drawString(36, H-45, ('前段' if i<30 else '后段') + f'  {rows[0][0]}:{rows[0][1]} 至 {rows[-1][0]}:{rows[-1][1]}')
        y = H-64
        for n, parts in enumerate(wrapped, 1):
            pdf.setFont('CJK', 7)
            pdf.setFillGray(.45)
            pdf.drawRightString(53, y, str(n))
            pdf.setFont('CJK', size)
            pdf.setFillGray(0)
            for part in parts:
                pdf.drawString(61, y, part)
                y -= leading
        assert y > 30
        pdf.showPage()
        evidence.append({'page':i+1, 'original_nonblank_lines':50, 'rendered_lines':count, 'font_size':size,
                         'mapping':[{'file':r[0], 'line':r[1], 'sha256':hashlib.sha256(r[2].encode()).hexdigest()} for r in rows]})
        text_pages.append(HEADER+f' 第{i+1}页\n'+'\n'.join(r[2] for r in rows))
    pdf.save()
    (OUT/'source-code.txt').write_text('\n\f\n'.join(text_pages)+'\n')
    assert len(PdfReader(out).pages) == 60
    return {'source_files':len(files), 'nonempty_source_lines':len(records),
            'selection':'Frozen committed web and minigame local TypeScript/TSX import dependencies. Entries first, remaining paths sorted. Blank lines omitted. First and last 1500 original nonblank lines; 50 original lines per submitted page, long lines only visually wrapped.',
            'files':[{'path':str(p.relative_to(SNAP)), 'sha256':digest(p)} for p in files],
            'excluded_generated_resource_files':[{'path':str(p.relative_to(SNAP)), 'sha256':digest(p)} for p in excluded],
            'pages':evidence}


def style_font(style, size, bold=False):
    family = 'Heiti SC' if bold else 'Songti SC'
    style.font.name = family
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    fonts = style.element.get_or_add_rPr().rFonts
    for key in list(fonts.attrib):
        if 'theme' in key.lower():
            del fonts.attrib[key]
    fonts.set(qn('w:eastAsia'), family)


def doc_base(kind):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin, s.bottom_margin = Inches(.65), Inches(.6)
    s.left_margin = s.right_margin = Inches(.7)
    s.header_distance = s.footer_distance = Inches(.25)
    for name, size, bold in [('Normal',11,False), ('Title',20,True), ('Heading 1',15,True), ('Heading 2',12,True), ('Caption',9,False)]:
        style_font(d.styles[name],size,bold)
    d.styles['Normal'].paragraph_format.space_after = Pt(7)
    d.styles['Normal'].paragraph_format.line_spacing = 1.18
    # The bundled default Word template carries a blue Title border.
    # Formal registration documents use a plain black title.
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    hp = s.header.paragraphs[0]
    hp.text = HEADER + '  ' + kind
    hp.style = d.styles['Caption']
    hp.paragraph_format.space_after = Pt(0)
    # Page numbering is at the upper right, alongside the software header.
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.35))
    hp.add_run('\t')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); hp._p.append(field)
    fp = s.footer.paragraphs[0]; fp.text = HEADER; fp.style = d.styles['Caption']
    d.core_properties.title = HEADER + ' ' + kind
    d.core_properties.author = ''
    return d


def paras(d, texts):
    for text in texts:
        d.add_paragraph(text)


def shot(d, filename, caption):
    p = d.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(OUT/'screenshots'/filename), width=Inches(6.8))
    p = d.add_paragraph(caption, 'Caption')
    p.paragraph_format.space_after = Pt(9)


def table(d, rows):
    t = d.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    for a,b in rows:
        cells=t.add_row().cells
        cells[0].text = a; cells[1].text = b
        cells[0].width=Inches(1.45);cells[1].width=Inches(5.25)
        for c in cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.size=Pt(10)
        trPr=t.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
    return t


def manual():
    d = doc_base('用户操作说明书')
    pages = [
      ('软件概述', '01-home.png', '图1 软件首页', [
        '本说明书介绍灰线像素前线游戏软件 V1.0 的启动、编队配置、对局操作、故事战役、卡牌收藏和声音设置，供用户了解软件功能与操作方法。',
        '软件采用横向像素战场与卡牌部署方式。玩家选择编队后与计算机对手进行即时策略对局，通过侦察、部署、支援和资源管理完成胜利目标。',
        '软件提供首页、故事战役、卡组、商店、设置和作战手册六个主要入口。故事战役包含四幕二十关，自由对战提供村镇、雨林、山地和沙漠四类地图。',
        '故事、人物、阵营与地图均为架空设定。软件名称和版本号用于登记材料的一致标识；界面显示的 v235 为该材料对应的内部构建编号。',
        '本说明书截图来自同一稳定构建的实际浏览器运行界面。小游戏端共享规则并使用 Canvas 界面与平台适配层；本文截图展示网页端操作。']),
      ('运行环境与启动', '05-settings.png', '图2 难度与声音设置', [
        '网页端在支持 JavaScript、Canvas 与本地存储的现代浏览器中运行，支持电脑鼠标键盘和手机触控操作。手机进入战场时采用横屏。',
        '打开软件首页后，先选择地图、昼间或夜间，再选择对手难度。标准对手携带25张牌，老练对手30张，精锐对手35张；对应回点速度为标准、加快15%和加快30%。',
        '进入卡组检查当前编队后，点击首页中央的开始游戏按钮进入自由对战。也可从故事战役进入已解锁任务。',
        '首次加载需要取得界面图片、字体和战场资源。页面正常显示后即可操作。浏览器限制音频自动播放时，点击界面或调整声音开关可触发音频初始化。',
        '地图、难度、卡组、收藏、战役进度和声音偏好保存在当前设备的本地存储中；清除浏览器站点数据、更换浏览器或设备不会自动同步这些记录。']),
      ('编队配置与卡牌详情', '02-deck.png', '图3 编组室与当前编队', [
        '点击卡组打开编组室，使用顶部下拉框切换编队。软件最多保存6套编队，支持重命名、另存为新卡组和删除编队。',
        '每套出战编队包含25张牌。同一种卡牌受拥有数量和携带上限约束；尚未获得的卡牌显示抽卡解锁，不能直接加入。',
        '通过步兵、载具、火炮、工事、航空、指令分类，费用筛选、编入状态筛选和搜索栏查找卡牌。卡面显示费用、名称、编制、属性与功能摘要。',
        '点击加号加入一张，点击减号减少一张。满编时选择新卡后按界面提示选择旧卡替换。右侧显示已编入卡牌、数量、平均费用与费用分布。',
        '长按卡面或点击详情按钮查看大图及完整说明。推荐编队提供配置参考，撤销可恢复最近一次修改；点击保存卡组保留编队，点击出战进入对局。']),
      ('故事战役与任务目标', '03-campaign.png', '图4 战役四幕与任务简报', [
        '点击故事战役进入断线之地。界面显示已完成关数、四幕进度、任务列表以及当前任务的地图、时段、简报、目标与已知部署。',
        '未解锁任务标记为待解锁。选择已解锁任务并点击进入任务，先阅读无线电简报，阅读期间战斗暂停，再进入任务战场。',
        '战役按任务配发借用编队和起始部队，起始指挥点由任务设定，因此可能不同于自由对战的零点开局。借用卡用于当前任务，不等同于永久收藏。',
        '任务分为坚守、夺取和反击等目标。坚守任务要求在规定时限内保护己方指挥部；夺取任务需占领指定目标点；反击任务按简报完成进攻要求。',
        '占领区域存在敌方争夺时，进度不能持续正常增长。用户应以战场任务提示为准安排部队。任务完成记录保存在当前设备，可返回菜单继续解锁后续关卡。']),
      ('自由对战与出牌操作', '07-battle.png', '图5 战场状态与战术手牌', [
        '自由对战的双方指挥部初始生命均为1000。摧毁敌方指挥部获胜；十分钟结束仍未决出胜负时比较双方基地剩余生命，生命相同时按规则结算平局。',
        '双方各使用自己的牌组。自由对战初始手牌6张、指挥点0点，基础每5.5秒恢复1点，初始上限10点；后勤、扩编和指令效果会改变恢复或上限。',
        '将手牌拖出底部区域后松手即可尝试使用，拖回手牌区域取消。普通地面部队从己方基地入场；烟幕、地雷、弹药空投和机降等卡牌按松手位置指定落点。',
        '出牌需要足够指挥点与合法目标。长按查看详情。点击牌堆或按R消耗2点抽一张牌，冷却1秒，手牌最多6张；无法抽牌时界面提示原因。',
        '左右拖动画面或操作小地图查看战线。电脑可用A/D、方向键或滚轮移动视野，数字键1至6选牌，回车使用，空格暂停。界面显示基地生命、时间、资源和战场通讯。']),
      ('部队协同与前线补给', '07-battle.png', '图6 部队部署后的战线', [
        '部队部署后按兵种和状态自主移动、观察、选择目标、射击、装填或执行支援。不同兵种具有独立的观察距离、射程、弹药、护甲与目标选择规则。',
        '观察范围影响战争迷雾与目标获取。地图会保留已观测敌军的接触信息；最后观测标记表示已知位置，并不持续揭示敌军的隐藏移动。',
        '点击己方小队可查看并下达伴随、据守、撤退、进攻或警戒等指令。战场还提供驻守、推进、奔跑、蹲行和卧倒等步兵行动选项。',
        '弹药与油料属于有限资源，血条下方细线显示弹药状态。弹药不足的可移动单位寻找补给，固定炮位等待前线补给；无弹药武器停止射击，车辆耗尽油料后不能继续移动。',
        '弹药空投需选择己方可见落点。补给箱、补给人员和补给车按库存补弹或补油。医疗人员救治伤兵，维修人员和车辆修复受损装甲；工事按各自建造时间与守军配置展开。']),
      ('卡牌收藏与军需商店', '04-shop.png', '图7 卡包与收藏列表', [
        '点击商店查看金币、卡包和已拥有卡牌。金币为游戏内资源；本文版本的商店操作不包含真实货币支付。',
        '点击开一包消耗100金币并获得5张卡牌；点击连开十包消耗1000金币并获得50张卡牌。界面展示本次结果和更新后的收藏数量。',
        '卡牌稀有度分为常规、精锐、王牌和传奇。达到可持有或携带数量限制的重复卡按相应规则转为金币；界面列出出率和溢出转换标准。',
        '收藏列表展示卡牌名称和拥有数量。点击卡牌可放大查看说明，取得的新卡可返回编组室配置到编队中。收藏与金币保存在当前设备。',
        '当前网页版本的看广告领金币按钮明确标记为模拟，仅用于本地体验。本文不将该入口描述为已接入抖音真实广告服务，也不将测试赠币视为正式运营充值。']),
      ('暂停声音与返回整备', '08-pause.png', '图8 对局暂停菜单', [
        '点击战场顶部的暂停作战或按空格打开暂停菜单。暂停后对局模拟停止推进，界面显示战场已暂停；点击继续作战恢复当前对局。',
        '暂停菜单可开关游戏声音并分别调整音乐与音效音量。设置保存到当前设备，首页设置页与战场使用同一声音偏好。',
        '点击重新整备或返回整备离开当前战场，回到菜单配置编队和下一局选项。离开对局不等同于保存该局的中途战斗状态。',
        '查看卡组可以了解对局编队情况。战场通讯与操作提示用于说明部署结果、抽牌条件和当前战线状态。',
        '音乐与音效署名入口列出外部音频的作者、来源、许可和使用修改信息。软件保留相应署名；登记软件程序不代表取得这些外部素材的独占权利。']),
      ('作战手册与软件结构', '06-guide.png', '图9 软件内置作战手册', [
        '点击作战手册查看架空设定、胜利条件、出牌抽牌、查看战场、机步协同、弹药补给和组建编队等说明。该入口用于在开始对局前或遇到操作疑问时查阅。',
        '软件采用TypeScript开发。战斗规则、卡牌数据、经济、地图、战役、收藏与存档模块构成共享逻辑；网页界面使用React，战场图形使用Canvas绘制。',
        '小游戏端通过Canvas界面、触摸路由和平台适配层实现菜单与战场交互，并适配图片、音频、本地存储等运行接口。Vite用于生成对应构建产物。',
        '程序维护对局时间、双方资源、牌库和手牌、部队及地图状态。玩家输入进入规则判定后更新状态，再由界面和绘制模块显示结果。',
        '本说明书对应已冻结的稳定源码和实际网页构建。小游戏工程仍需配置真实平台应用标识并完成平台联调；其工程存在不等于已通过抖音审核或正式上架。'])
    ]
    for i,(title,img,caption,body) in enumerate(pages):
        if i==0:
            d.add_paragraph(TITLE, 'Title')
            d.add_paragraph(VERSION + ' 用户操作说明书')
        heading = d.add_paragraph(f'{i+1} {title}', 'Heading 1')
        if i:
            heading.paragraph_format.page_break_before = True
        # Put introductory text before the illustration so every page gives
        # context immediately, followed by concrete operations.
        d.add_paragraph(body[0])
        shot(d,img,caption)
        paras(d,body[1:])
    d.save(OUT/'02-user-manual.docx')
    (OUT/'manual-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2))


def application(source):
    d = doc_base('申请填报资料与材料清单')
    d.add_paragraph('灰线软件著作权登记填报资料', 'Title')
    paras(d,['本文件是填报资料草稿和办理清单，不是中国版权保护中心生成的正式申请表或签章页。名称与版本暂沿用现有材料，待申请人确认。著作权人、真实开发完成日期、发表及权属信息未核实前不能提交。'])
    d.add_paragraph('1 可直接核对的软件资料','Heading 1')
    table(d,[('软件全称',TITLE),('软件简称','像素前线'),('版本号',VERSION),('软件类型','游戏软件'),('编程语言','TypeScript 含TSX界面文件'),('源程序量',f"{source['nonempty_source_lines']}个非空源程序行，{source['source_files']}个入口依赖文件。仅统计本项目已冻结的网页与小游戏逻辑，不计第三方库、生成图片数据和构建产物。"),('软件主要功能','提供像素即时策略对局、25张卡组编配、四幕二十关战役、卡牌收藏、游戏内金币卡包、战争迷雾、部队协同及有限补给，支持本地保存与声音设置。'),('开发目的','实现面向电脑与移动设备的单人即时策略卡牌游戏，支持玩家编队并与计算机对手对战。'),('技术特点','TypeScript共享游戏规则，React网页界面与Canvas战场绘制，小游戏平台适配，本地存档、像素动画及分模块资源加载。')])
    d.add_page_break()
    d.add_paragraph('2 开发与运行信息','Heading 1')
    table(d,[('开发硬件环境','待申请人填写实际开发设备、处理器与内存；不以当前整理材料的机器替代开发事实。'),('开发操作系统','待申请人确认实际开发环境。'),('开发工具','项目可确认：TypeScript、Vite、React。实际使用的编辑器由申请人确认。'),('运行硬件环境','支持图形界面、浏览器与触控或鼠标键盘输入的电脑或智能手机；具体最低硬件规格未测定。'),('运行软件环境','网页端：支持JavaScript、Canvas、本地存储的现代浏览器。小游戏端：抖音小游戏接口适配工程，真实平台应用标识和实机联调尚待完成。'),('对应源码构建','v235 稳定提交 67c4c4b6a4b4eadb09fe8e85db30ae0013b269e3；不包含当前工作区尚未提交的战斗修复。')])
    d.add_paragraph('3 必须由申请人提供的真实信息','Heading 1')
    table(d,[('主体类型与全称','待填：个人姓名或公司全称，须与实名认证及证件一致。'),('开发完成日期','待填：本次申请版本实际开发完成日期。整理材料日期不是开发完成日期。'),('发表情况','待填：是否已发表，以及首次公开日期、地点、方式。项目存在公开网页记录，未上架抖音不代表未发表。'),('开发与权属关系','待填：独立、合作、委托或职务开发事实；原始取得、受让等取得方式及拟申请权利范围。')])
    d.add_page_break()
    d.add_paragraph('4 证件与补充证明文件','Heading 1')
    paras(d,[
      '个人申请：按官方系统要求准备有效身份证明并完成个人实名认证。公司申请：准备营业执照等法人身份证明并完成机构实名认证。证件号码、证件图片及私人联系方式直接在官方系统填写或上传，不写入公开项目仓库。',
      '联系人、联系电话、邮箱、通信地址和邮编需由申请人核实。若存在合作开发、委托开发或项目任务书，准备真实合同、任务书或权属约定；受让、继承或经许可修改的软件，准备相应证明。没有该情形时不制作虚构协议。',
      '应核对是否属于职务开发、是否使用公司的开发任务和资源、是否有其他开发参与者及权属约定。不能仅凭由某人操作电脑就认定其拥有全部软件著作权。'])
    d.add_paragraph('5 已整理的鉴别材料','Heading 1')
    table(d,[('源程序鉴别材料','01-source-code.pdf；前后各连续30页，共60页，每页50个原始非空源代码行，长行仅作版面折行。页眉名称与版本一致，末页保留原程序末尾。'),('用户操作说明书','02-user-manual.pdf；提供真实网页截图及主要模块操作说明，整份提交。可编辑版本为02-user-manual.docx。'),('填报资料','03-application-and-checklist.pdf及可编辑Word文件；辅助填写，不替代官方申请表。'),('追溯附件','copyright-evidence.json、snapshot-manifest.json和source-code.txt，用于核对源码、页面与文件哈希，不作为默认上传附件。')])
    d.add_page_break()
    d.add_paragraph('6 第三方素材与权利说明','Heading 1')
    paras(d,[
      '源程序材料不打包node_modules、React等第三方库源码，不删除第三方许可声明来制造原创。源码内容来自稳定提交，不新增模拟函数来凑行数。',
      '项目内音频署名文件列有CC0音频，以及Michel Baradari的爆炸音效和jakobthiesen及AntumDeluge的循环飞行音频等CC BY 3.0素材。应保留作者、许可链接及修改说明。该清单依据项目现有署名，不构成对所有素材许可完整性的法律确认。',
      '项目包含Galmuri、Ark Pixel和Cubic 11字体许可文件。对应字体及第三方音频不申报为申请人自行创作的独占素材。',
      '部分像素精灵和卡牌插画通过生成工具制作。申请人应保留制作记录、使用条款与后续人工编辑记录；软件著作权登记材料不应顺带主张全部生成图像都已取得独占美术著作权。',
      '本次申请对象为软件程序及相应文档。游戏名称、商标、图像作品权属、出版或平台运营资质属于不同事项；本材料包不代表这些事项已获批准。'])
    d.add_paragraph('7 官方办理顺序','Heading 1')
    paras(d,[
      '第一步：从中国版权保护中心官网进入软件登记，登录并完成主体实名认证。',
      '第二步：确认软件名称、版本、主体、完成日期、首次发表和权属关系，按本草稿录入真实信息。',
      '第三步：按系统要求上传源程序和说明书，以及适用的身份证明、合同或权属证明。单文件格式和大小限制以当次系统提示为准。',
      '第四步：由官方系统生成申请确认签章页，核对内容后由申请人真实签名或公司盖章，再上传扫描件。不要使用本草稿冒充正式表格，不制作或粘贴虚假签章。',
      '第五步：最终核对并提交，保存申请流水号和回执，关注受理或补正通知；收到补正后按通知要求处理。提交成功、正式受理和取得证书是不同状态。'])
    d.add_page_break()
    d.add_paragraph('8 依据与当前办理状态','Heading 1')
    paras(d,[
      '《计算机软件著作权登记办法》第九至十一条规定申请材料、程序与文档鉴别材料以及身份证明等要求。一般交存为程序和一种文档的前后各连续30页；不足60页时提交全部，除特定情况外程序每页不少于50行、文档每页不少于30行。含图文档版式还应遵从登记系统的具体提示。',
      '中国版权保护中心2023年网上办理通知说明，软件登记通过系统在线填报并上传申请文件，无需再邮寄新申请纸介质材料。材料审查可能要求补正。费用、文件限制和具体办理安排以当前官方系统为准。',
      '截至2026年10月14日，本材料包已完成独立于申请主体的准备。主体、名称版本、完成日期、发表及权属信息仍待申请人确认；未进行实名认证、未填写签章声明、未向官方系统上传或提交本次申请，尚无本次申请流水号或登记证书。',
      '本次采用已提交的v235作为稳定基线，未将工作区正在修改的战斗代码写入材料。若实际申报对象换成后续完成版本，应重新冻结对应源码与界面，并同步更新名称、版本、说明书、统计和鉴别材料。'])
    d.add_paragraph('官方参考入口','Heading 2')
    paras(d,[LAW, NOTICE, PORTAL])
    d.save(OUT/'03-application-and-checklist.docx')


def main():
    OUT.mkdir(exist_ok=True)
    source = source_pdf()
    manual()
    application(source)
    snapshot = json.loads((OUT/'snapshot-manifest.json').read_text())
    evidence = {'prepared_at':'2026-10-14', 'software_name':TITLE,'version':VERSION,
      'name_and_version_status':'Awaiting applicant confirmation', 'source_commit':snapshot['commit'],
      'source':source, 'screenshots': [{'file':p.name,'sha256':digest(p)} for p in sorted((OUT/'screenshots').glob('*.png'))],
      'screenshot_origin':'Actual browser UI of committed v235 web build served locally from docs on port 8240; no simulated Douyin screenshots.',
      'status':'Draft materials prepared. Applicant identity, dates, publication and ownership not yet confirmed; not submitted.',
      'official_references':[LAW, NOTICE, PORTAL]}
    (OUT/'copyright-evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'source_files':source['source_files'],'source_lines':source['nonempty_source_lines'],
      'source_pages':60,'minimum_font_size':min(p['font_size'] for p in source['pages']),
      'output':str(OUT)},ensure_ascii=False))


if __name__ == '__main__':
    main()
